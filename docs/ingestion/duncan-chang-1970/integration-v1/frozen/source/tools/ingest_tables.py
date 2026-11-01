"""Reviewed worksheet blocks, retaining native units and independent coordinates.

These are measurement series, not normalized triaxial experiments. A worksheet
may contain several independent curves: its registry configuration splits them
into separate records instead of joining by Excel row number.
"""
from __future__ import annotations

import hashlib
from datetime import date, datetime, time
import io
import re
import zipfile

from atlas.public_data import json_bytes


def source_cell(value):
    return value.isoformat() if isinstance(value, (date, datetime, time)) else value


def workbook_context(workbook):
    return {name: [[i+1, *[source_cell(v) for v in row]]
                   for i,row in enumerate(workbook[name].iter_rows(values_only=True))
                   if any(v is not None for v in row)]
            for name in ('Basic Properties', 'References') if name in workbook.sheetnames}


def numeric(value):
    import math
    if value is None or isinstance(value, bool):
        return None
    try:
        result = float(value.replace(',', '.').strip() if isinstance(value, str) else value)
        return result if math.isfinite(result) else None
    except (TypeError, ValueError):
        return None


def source_series(source, file, sheet, config, *, archive_sha256, member_sha256=None,
                  workbook_info=None):
    """Extract an explicitly reviewed, coordinate-aligned source block."""
    start = config['data_start']
    stop = config.get('data_end', len(sheet))
    selected = config['columns']
    if len(set(selected)) != len(selected) or len(selected) < 2:
        raise ValueError('A source series needs distinct reviewed columns')
    headers = config['header_rows']
    units_row = config.get('units_row')
    columns = []
    for col in selected:
        labels = [str(sheet[r - 1][col - 1]).strip() for r in headers
                  if col <= len(sheet[r - 1]) and sheet[r - 1][col - 1] is not None]
        if not labels:
            raise ValueError('Reviewed source column has no header')
        unit = config.get('units', {}).get(str(col))
        if unit is None and units_row:
            unit = sheet[units_row - 1][col - 1]
        columns.append(dict(key='c' + str(col), name=' · '.join(dict.fromkeys(labels)),
                            unit=str(unit) if unit is not None else 'not specified', source_column=col))
    native, rows, row_numbers = [], [], []
    for row_number, values in enumerate(sheet[start - 1:stop], start):
        if config.get('data_rows') and row_number not in config['data_rows']:
            continue
        raw = [values[col - 1] if col <= len(values) else None for col in selected]
        if not any(v is not None for v in raw):
            continue
        native.append([row_number, *raw])
        normalized = {c['key']: n for c, value in zip(columns, raw) if (n := numeric(value)) is not None}
        if normalized:
            rows.append(normalized)
            row_numbers.append(row_number)
    if not rows:
        raise ValueError('Selected source block contains no numeric observations')
    # All source text and nonnumeric cells are retained, including loading stage
    # strings, native time notation and cached spreadsheet formula errors.
    metadata = [[i + 1, *r] for i, r in enumerate(sheet[:start - 1]) if any(v is not None for v in r)]
    provenance = {k: source[k] for k in ('source_url', 'license', 'license_url', 'attribution')}
    provenance.update(source_id=source['id'], archive=file['filename'], archive_sha256=archive_sha256,
                      member=config['sheet'], member_sha256=member_sha256 or hashlib.sha256(json_bytes(native)).hexdigest(),
                      member_hash_scope='Original workbook bytes' if member_sha256 else 'Selected native cells and Excel row numbers, canonical JSON',
                      source_rows=row_numbers, source_metadata=metadata, workbook_info=workbook_info or [],
                      original_references=config.get('original_references', []),
                      transformations=['Source numeric values retained in native units and order; decimal-comma numeric strings parsed.',
                                       'Nonnumeric and missing cells omitted from numeric channels and retained in source_table.',
                                       'No interpolation, stress inference, sign conversion or cross-curve coordinate join.'],
                      processor='tools/ingest_tables.py:source-series-v1')
    warnings = ['Source units and signs are retained as reported; this series is not a normalized triaxial model overlay.',
                *config.get('warnings', [])]
    return dict(schema_version='0.1', kind=config.get('kind', 'measurement-series'),
                id=config['id'], material=config.get('material', file['material']), protocol=config['protocol'],
                conventions=dict(stress='source-defined', strain='source-defined',
                                 stress_unit='per-channel', strain_unit='per-channel'),
                initial={}, columns=columns, rows=rows, provenance=provenance, warnings=warnings,
                source_plot=dict(x='c' + str(config.get('x', selected[0])),
                                 y='c' + str(config.get('y', selected[1]))),
                source_table=dict(columns=['excel_row', *[c['name'] for c in columns]],
                                  units=['1', *[c['unit'] for c in columns]], values=native))


def read_workbook_tables(source, file, raw, archive_sha256, *, member_sha256=None):
    import openpyxl
    workbook = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    records = []
    try:
        info = [[source_cell(v) for v in row] for row in workbook['info'].iter_rows(values_only=True)] if 'info' in workbook.sheetnames else []
        context = workbook_context(workbook)
        configs = file.get('tables', [])
        if file.get('all_specimens'):
            # Finnish raw workbooks: one specimen per sheet; columns A:M are
            # aligned observations. Later side columns hold specimen constants.
            configs = []
            for sheet in workbook:
                if sheet.title == 'info':
                    continue
                header = next(sheet.iter_rows(max_row=1, values_only=True))
                count = min(13, len(header))
                cols = [i + 1 for i, name in enumerate(header[:count]) if name is not None and
                        isinstance(name, str) and name not in ('stage',)]
                units = {'stress': 'kPa', 'height': 'mm', 'strain_pros': '%', 'strain': '1',
                         'e': '1', 'e0': '1', 'v': '1', 'ln_sigma': 'ln(source kPa)',
                         'log_sigma': 'log10(source kPa)', 'stress_mean': 'kPa',
                         'time': 'native time notation; see data notes', 'modulus': 'source-defined'}
                stress = header.index('stress') + 1
                y = header.index('e') + 1
                configs.append(dict(id=file['id_prefix'] + '-' + sheet.title, sheet=sheet.title,
                                    protocol='oedometer', columns=cols, header_rows=[1], data_start=2,
                                    units={str(i): units.get(header[i - 1], 'not specified') for i in cols},
                                    x=stress, y=y, warnings=[
                                        'Stress is effective vertical stress in kPa according to the workbook info sheet.',
                                        'Time notation includes days/hours/minutes or native numbers; it is not interpreted as cumulative elapsed time.',
                                        'Some early observations have missing stress. Loading/unloading stage strings remain in the native table.']))
                # Keep stage and other nonnumeric source columns in the archive.
                # They are not graph axes, but their original row alignment matters.
        for config in configs:
            sheet = [[source_cell(v) for v in row] for row in workbook[config['sheet']].iter_rows(values_only=True)]
            record = source_series(source, file, sheet, config, archive_sha256=archive_sha256,
                                   member_sha256=member_sha256, workbook_info=info)
            if context:
                record['provenance']['workbook_context'] = context
            if file.get('all_specimens') or config.get('retain_original_rows'):
                record['source_table']['original_rows'] = [[i + 1, *r] for i, r in enumerate(sheet)
                                                         if any(v is not None for v in r) and
                                                         (not config.get('data_rows') or i == 0 or i + 1 in config['data_rows'])]
            records.append(record)
    finally:
        workbook.close()
    return records


def read_tables(source, raw_dir):
    records = []
    for file in source['files']:
        raw = (raw_dir / file['filename']).read_bytes()
        digest = hashlib.sha256(raw).hexdigest()
        if digest != file['sha256']:
            raise ValueError('Source checksum mismatch: ' + file['filename'])
        if file.get('workbooks'):
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                for member in file['workbooks']:
                    data = archive.read(member['member'])
                    sha = hashlib.sha256(data).hexdigest()
                    if sha != member['sha256']:
                        raise ValueError('Workbook member checksum mismatch')
                    child = dict(member, filename=file['filename'])
                    records.extend(read_workbook_tables(source, child, data, digest, member_sha256=sha))
        elif file.get('tables'):
            records.extend(read_workbook_tables(source, file, raw, digest))
    return records
